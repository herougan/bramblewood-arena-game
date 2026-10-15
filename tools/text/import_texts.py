#!/usr/bin/env python3
"""Apply the "New text" column of an edited docs/game-texts.xlsx back to the game (2026-10-15).

    python3 tools/text/import_texts.py EDITED.xlsx [--dry-run] [--root DIR]

  * keyed rows (toast., dialog., coach., tutorial., results.): the value in bramblewood-texts.js is replaced. One-to-one.
  * src: rows: the literal is found again in arena_app.js / arena_template.html by its key (or, failing that, by its exact
    current text when that is unique) and rewritten in place, every place the row covers.
A row is skipped, and listed in the report, when its key is unknown, the code's text changed since the export (the sheet's
"Current text" no longer matches), the text can't be found or matches more than one row, or the new text adds a {placeholder}
the code doesn't have. Rebuild afterwards (python3 assemble_arena.py && mv bramblewood-arena.html index.html).
Exit code 0 when everything applied, 2 when some rows were skipped.
"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import textlib as T


def read_rows(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb['Game texts'] if 'Game texts' in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows[:10]) if r and any((str(c or '').strip().lower() == 'key') for c in r))
    hdr = [str(c or '').strip().lower() for c in rows[hdr_i]]
    col = {name: hdr.index(name) for name in ('key', 'current text', 'new text') if name in hdr}
    out = []
    for n, r in enumerate(rows[hdr_i + 1:], start=hdr_i + 2):
        key = r[col['key']] if col['key'] < len(r) else None
        if not key or str(key).startswith('('): continue
        new = r[col['new text']] if col['new text'] < len(r) else None
        cur = r[col['current text']] if 'current text' in col and col['current text'] < len(r) else None
        if new is None or str(new) == '': continue
        out.append({'row': n, 'key': str(key).strip(), 'current': '' if cur is None else str(cur), 'new': str(new)})
    return out


def norm(s):
    return (s or '').replace('\r\n', '\n').replace('\r', '\n')


def apply(rows, root=T.ROOT, dry=False):
    report = {'applied': [], 'unchanged': [], 'skipped': [], 'warnings': []}
    table = T.load_table(root)
    table_changed = False
    src_rows = []
    for r in rows:
        new, cur, key = norm(r['new']), norm(r['current']), r['key']
        if key.startswith('src:'):
            src_rows.append(r); continue
        if key not in table:
            report['skipped'].append((r, 'unknown key (not in bramblewood-texts.js)')); continue
        if new == table[key]:
            report['unchanged'].append((r, 'already says that')); continue
        if cur and cur != table[key]:
            report['skipped'].append((r, 'the game text changed since this sheet was exported; export again')); continue
        ok, msg = T.placeholders_ok(table[key], new, numbered=False)
        if not ok:
            report['skipped'].append((r, msg)); continue
        if msg: report['warnings'].append((r, msg))
        table[key] = new; table_changed = True
        report['applied'].append((r, 'bramblewood-texts.js'))
    edits = {}            # file -> [(start, end, raw, row)]
    srcs = {}
    if src_rows:
        units = T.extract_units(root)
        by_key = {u.key: u for u in units}
        by_text = {}
        for u in units: by_text.setdefault(u.text, []).append(u)
        for r in src_rows:
            new, cur = norm(r['new']), norm(r['current'])
            u = by_key.get(r['key'])
            if u is None:
                cands = by_text.get(cur, [])
                if len(cands) == 1: u = cands[0]; report['warnings'].append((r, 'matched by its text (the key moved): ' + u.key))
                elif len(cands) > 1:
                    report['skipped'].append((r, 'ambiguous: this text is in %d different places; export again' % len(cands))); continue
                else:
                    report['skipped'].append((r, 'not found in the code any more; export again')); continue
            if cur and u.text != cur:
                report['skipped'].append((r, 'the game text changed since this sheet was exported; export again')); continue
            if new == u.text:
                report['unchanged'].append((r, 'already says that')); continue
            ok, msg = T.placeholders_ok(u.text, new)
            if not ok:
                report['skipped'].append((r, msg)); continue
            if msg: report['warnings'].append((r, msg))
            for rel, run in u.runs:
                if rel not in srcs: srcs[rel] = open(os.path.join(root, rel), encoding='utf-8').read()
                quote = '' if run.lit.kind == 'h' else srcs[rel][run.lit.start]
                raw = new if run.lit.kind == 'h' else T.rebuild_raw(run, new, quote)
                if run.attr and '"' in new:
                    raw = raw.replace('"', '&quot;')
                edits.setdefault(rel, []).append((run.start, run.end, raw, r))
            report['applied'].append((r, ', '.join(sorted({rel for rel, _ in u.runs}))))
    # a row whose text sits inside another edited row's {n} part can't go in the same pass: keep the outer one, report the inner
    dropped = set()
    for rel, lst in edits.items():
        lst.sort(key=lambda e: (e[0], -(e[1] - e[0])))
        end = -1; outer = None
        for a, b, raw, r in lst:
            if a < end and r is not outer:
                dropped.add(id(r))
            elif b > end:
                end, outer = b, r
    if dropped:
        keep = []
        for r, where in report['applied']:
            if id(r) in dropped: report['skipped'].append((r, 'sits inside another row edited in this sheet: run the import again to apply it'))
            else: keep.append((r, where))
        report['applied'] = keep
    for rel, lst in edits.items():
        p = os.path.join(root, rel)
        src = open(p, encoding='utf-8').read()
        for a, b, raw, r in sorted([e for e in lst if id(e[3]) not in dropped], key=lambda e: e[0], reverse=True):
            src = src[:a] + raw + src[b:]
        if not dry:
            open(p, 'w', encoding='utf-8').write(src)
    if table_changed and not dry:
        T.save_table(table, root)
    return report


def print_report(rep, dry):
    tag = ' (dry run: nothing written)' if dry else ''
    print('import%s: %d applied, %d unchanged, %d skipped' % (tag, len(rep['applied']), len(rep['unchanged']), len(rep['skipped'])))
    for r, where in rep['applied']:
        print('  ✓ row %d  %s  -> %s' % (r['row'], r['key'], where))
    for r, why in rep['warnings']:
        print('  ! row %d  %s: %s' % (r['row'], r['key'], why))
    for r, why in rep['skipped']:
        print('  ✗ row %d  %s: %s' % (r['row'], r['key'], why))


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    dry = '--dry-run' in sys.argv
    root = T.ROOT
    if '--root' in sys.argv: root = os.path.abspath(sys.argv[sys.argv.index('--root') + 1]); args = [a for a in args if os.path.abspath(a) != root]
    if not args: raise SystemExit(__doc__)
    rows = read_rows(args[0])
    rep = apply(rows, root, dry)
    print_report(rep, dry)
    if not dry and rep['applied']:
        print('Now rebuild: python3 assemble_arena.py && mv bramblewood-arena.html index.html')
    sys.exit(2 if rep['skipped'] else 0)


if __name__ == '__main__':
    main()
