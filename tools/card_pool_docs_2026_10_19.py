#!/usr/bin/env python3
"""docs/card-pool-2026-10-19.md + .xlsx from docs/balance/card-pool-2026-10-19.json (written by
`node tools/card_pool_2026_10_19.js write`). 2026-10-19 (package AN): the candidate pool, one table / sheet per map."""
import json, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
rows = json.load(open(os.path.join(ROOT, 'docs/balance/card-pool-2026-10-19.json'), encoding='utf-8'))
import subprocess  # the map names / icons, read the way the tests read CONQUEST_MAPS
maps = {m['id']: (m['name'], m.get('icon', '')) for m in json.loads(subprocess.check_output(['node', '-e',
    "const s=require('fs').readFileSync(process.argv[1],'utf8'),i=s.indexOf('const CONQUEST_MAPS = ['),j=s.indexOf('\\n];',i);"
    "console.log(JSON.stringify(eval(s.slice(i+'const CONQUEST_MAPS = '.length,j+2)).map(m=>({id:m.id,name:m.name,icon:m.icon}))))",
    os.path.join(ROOT, 'arena_app.js')]))}
ORDER = [('m1', 0), ('mf', 1), ('mb', 2), ('mc', 3), ('m2', 4), ('mw', 5), ('m8', 6), ('mu', 7), ('m4', 8), ('mk', 9), ('m7', 10), ('m10', 11), ('mh', 12), ('mt', 16), ('m6', 17), ('mi', 17)]
RN = {'common': 'Common', 'uncommon': 'Uncommon', 'rare': 'Rare', 'veryrare': 'Very Rare'}

def label(mid, n):
    nm, ic = maps.get(mid, (mid, ''))
    return f"Map {n} · {ic} {nm} ({mid})"

md = ["# Candidate card pool (2026-10-19, package AN)", "",
      "Owner 23:04: *\"More thematic animals — I'll pick and choose when balancing for a map (5–10 new cards per map — from skirmish-only rewards and pack), like pin cushions (uni = sea urchins) and sea snakes. (Temperate water-going mammals?) (Do the polar ones in advance — walruses and stuff.)\"* 23:32: *\"We don't have a really scary spider. Not trapdoor or cute.\"*", "",
      f"**{len(rows)} candidate cards**: 5–7 for each of Maps 0–12, a polar set of {sum(1 for r in rows if r['map'] in ('m6', 'mi', 'mt'))} for the taiga and tundra (Maps 16–17), the owner's sea urchin and sea snakes, the temperate water mammals and one frightening spider.", "",
      "## How the pool works", "",
      "- Each card is in `canonical/cards.json` with `\"pool\": true`, `\"poolMap\": \"<map id>\"` and `\"poolSlot\": \"skirmish\" | \"pack\"` (the slot the card was made for). It has no `source`, is locked, and has a pool number 9001+.",
      "- **Players never see a pool card.** `assemble_arena.py` leaves it out of the game's card list (so no pack, reward, Codex / Bestiary, CPU deck, ghost deck, egg, auto-battle pool or card list has it), and the card hash ignores it (`bramblewood-integrity.js`), so the server's list and the page still agree. The test harness, the card matrix, `tools/lib_bramble.js` (deck tools) and `tools/card_balance.js`'s full report skip it too.",
      "- **Admin → Manage Cards → \"Pool cards (N)\"** lists them by map (Map 0 first), each tile tagged **Pool · skirmish reward** or **Pool · pack**. Hovering shows its note (idea, budget, win rate). A click opens it in the card editor as a new card, its map already filled in under From: pick the skirmish node (or Pack) and Save. Saving makes it a live card with the next live card number, and it drops off the pool list.",
      "- To re-measure: `node tools/card_balance.js 120 --pool` (or `--only=<ids>`) writes `docs/balance/card-balance-partial.json`. `node tools/card_pool_2026_10_19.js tune` re-tunes the specs, `write` puts them back in `cards.json`, and this script rebuilds this page and the .xlsx.", "",
      "**Rules every card meets** (`tests/card-pool-2026-10-19.js`): it uses existing skills only (Tide, never the retired Swimmer). Its rarity is within its map's skirmish road cap (Rare on Maps 0–12, Very Rare on 16–17). It passes Budget v3 at Level 1, with N, S and both rules in its creatorNote. Its measured win rate (120 games, the `card_balance.js` yardstick) is inside its rarity band (Common 35–60%, Uncommon 45–68%, Rare 52–75%, Very Rare 56–80%). Its emoji is in all three shipped sets. Its flavour is under 20 words. Its name and flavour are in all 12 language packs.", "",
      "## The temperate water mammals", "",
      "Yes, and they fit the river and coast maps well. The game already had the Otter faction, the Beaver Builder, the Harbor Seal, the Walrus Icebreaker, the Leopard Seal and the Platypus. New in the pool:", "",
      "- **Water Vole** (Map 0, the river bank), **River Mink** (Map 0, a mustelid that is not an Otter-faction card), **Muskrat Lodgekeeper**, **Water Shrew** (venomous bite) and **Pyrenean Desman** (Map 4, Sunken Hollow), **Grey Seal Bull** (Map 2, Pebble Beach). The polar set adds the **Harp Seal Pup**, the **Ringed Seal** and the **Beluga**, and Map 9 has the **Mangrove Manatee** (warm water).",
      "- Not added: **Sea Otter** (it would read as an Otter-faction card), **Beaver** (Beaver Builder exists), **Harbour Seal** (Harbor Seal exists), **Eurasian Otter** (the Otter faction).", "",
      "## The scary spider (owner request)", ""]
sp = next((r for r in rows if r['id'] == 'giant-cave-huntsman'), None)
if sp:
    md += [f"**{sp['icon']} {sp['name']}**, Map 8 Caves & Alcoves (dark), {RN[sp['rarity']]} (the map's cap), Lumber {sp['lumber']}, Wait {sp['wait']}, {sp['attack']}/{sp['health']}: {sp['skills']}. Measured {sp['winRate']}% (Rare band {sp['band']}%).", "",
           f"> *{sp['flavor']}*", "", sp['idea'], "", f"{sp['art']}", ""]
md += ["## The polar set", "", "Every polar card resists Cold. The Walrus Icebreaker, Narwhal Lancer, Snow Ermine, Arctic Fox Raider, Snow Owl Sentinel and Tundra Musk Ox already exist, so they are not repeated here.", ""]
md += [f"- {r['icon']} **{r['name']}** ({label(r['map'], 17 if r['map'] != 'mt' else 16)}): {RN[r['rarity']]}, {r['attack']}/{r['health']}, {r['skills']}" for r in rows if r['map'] in ('m6', 'mi', 'mt')]
md += ["", "## By map", ""]
hdr = "| Card | Slot | Rarity | Lumber | Wait | A/H | Skills | N | S | Win rate | Idea |"
for mid, n in ORDER:
    rs = [r for r in rows if r['map'] == mid]
    if not rs: continue
    md += [f"### {label(mid, n)} ({len(rs)})", "", hdr, "|" + "---|" * 11]
    md += [f"| {r['icon']} {r['name']} | {r['slot']} | {RN[r['rarity']]} | {r['lumber']} | {r['wait']} | {r['attack']}/{r['health']} | {r['skills']}{' · resists ' + r['resist'] if r['resist'] else ''} | {r['N']} | {r['S']} | {r['winRate']}% ({r['band']}) | {r['idea'].split(': ')[0] if len(r['idea']) > 140 else r['idea']} |" for r in rs]
    md += [""]
md += ["## Decisions for the owner", "",
       "- **Placement:** each card names a map and the slot it was made for (skirmish reward or pack). Nothing is placed: you choose 5–10 per map when you balance it.",
       "- **Stand-in emoji** where none exists: 🌰 for the Pincushion Urchin, 🦭 for the Mangrove Manatee, 🦜 (sea parrot) for the Puffin, 🦤 for the Bittern, 🐗 for the Truffle Piglet (no pig emoji in the sets).",
       "- **Horseshoe Crab** is typed Arachnid (it is a chelicerate, nearer spiders than crabs), so Arachnid synergies pick it up.",
       "- Win rates sit anywhere inside their band, some near the low edge. When you place a card, re-measure it against the map's decks.", ""]
open(os.path.join(ROOT, 'docs/card-pool-2026-10-19.md'), 'w', encoding='utf-8').write('\n'.join(md))

wb = Workbook(); wb.remove(wb.active)
cols = ['Card', 'Emoji', 'Id', 'Slot', 'Rarity', 'Lumber', 'Wait', 'Attack', 'Health', 'Skills', 'Resist', 'N', 'S', '2A+H+S', 'A*H+S', '(N/3)^2', 'Win rate %', 'Band %', 'Idea', 'Flavour', 'Art prompt']
def sheet(title, rs):
    ws = wb.create_sheet(title[:31]); ws.append(cols)
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='6A4BB0')
    for r in rs: ws.append([r['name'], r['icon'], r['id'], r['slot'], RN[r['rarity']], r['lumber'], r['wait'], r['attack'], r['health'], r['skills'], r['resist'], r['N'], r['S'], r['x1'], r['x2'], r['N2'], r['winRate'], r['band'], r['idea'], r['flavor'], r['art']])
    for i, w in enumerate([24, 6, 24, 9, 10, 7, 6, 7, 7, 34, 8, 7, 6, 8, 8, 8, 9, 8, 70, 60, 40]): ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(vertical='top', wrap_text=c.column in (10, 19, 20, 21))
    ws.freeze_panes = 'B2'
sheet('All', [r for mid, n in ORDER for r in rows if r['map'] == mid])
for mid, n in ORDER:
    rs = [r for r in rows if r['map'] == mid]
    if rs: sheet(f"Map {n} {mid}", rs)
wb.save(os.path.join(ROOT, 'docs/card-pool-2026-10-19.xlsx'))
print('wrote docs/card-pool-2026-10-19.md and .xlsx:', len(rows), 'cards')
