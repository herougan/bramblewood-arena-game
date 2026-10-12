#!/usr/bin/env python3
"""Card, map and effect workbook (2026-10-12, user: "Generate an csv with card names, attack, defence, abilities, type(s)
and where they appear, then on the 2nd sheet, the map list and submaps under them. On the third sheet, all effects and
when they are generated ... note about when the card packs appear as well").

    python3 tools/export_gamedata.py OUT_DIR      # snapshot from the built game (needs the local server on :8765)
    python3 tools/export_workbook.py OUT_DIR/gamedata.json OUT.xlsx

Writes OUT.xlsx (sheets: Cards, Maps, Effects) and one CSV per sheet next to it (OUT-cards.csv, OUT-maps.csv,
OUT-effects.csv). The card-pack timeline sits under the map table on the Maps sheet and in OUT-packs.csv.
"""
import csv, json, os, re, sys
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

data = json.load(open(sys.argv[1]))
out = sys.argv[2]
base = os.path.splitext(out)[0]

RLABEL = {'starter': 'Starter', 'common': 'Common', 'uncommon': 'Uncommon', 'rare': 'Rare', 'veryrare': 'Very Rare',
          'superrare': 'Super Rare', 'epic': 'Epic', 'heroic': 'Heroic', 'unique': 'Unique', 'legendary': 'Legendary',
          'mythic': 'Mythic', 'ancient': 'Ancient', 'quest': 'Quest', 'questunique': 'Quest-Unique', 'special': 'Special',
          'devlegendary': 'Dev-Legendary', 'devancient': 'Dev-Ancient', 'eventlegendary': 'Event-Legendary',
          'eventrare': 'Event-Rare', 'divine': 'Divine', 'void': 'Void', 'daemonic': 'Daemonic'}
KIND = {'unit': 'Unit', 'spell': 'Spell', 'equipment': 'Equipment', 'structure': 'Structure', 'materia': 'Materia', 'field': 'Field'}
PACK_POOL = {1: 'Basic packs (Sprout Pouch, Acorn Chest, Golden Bramble Case)', 11: 'Basic packs (Factions pool)', 99: 'Mystery Booster'}

# ---------- maps: numbering, labels, who fights where ----------
maps = [m for m in data['maps'] if not (m['id'].startswith('r') and not m['nodes'])]
by_id = {m['id']: m for m in maps}
mains = [by_id[i] for i in data['mainOrder'] if i in by_id and not by_id[i]['sub'] and by_id[i]['no'] is not None and by_id[i]['no'] <= 20]
main_ids = {m['id'] for m in mains}
subs_of = defaultdict(list)
for m in maps:
    if m['sub'] and m['parent'] in main_ids:
        subs_of[m['parent']].append(m)

def map_label(m):
    if m['sub']:
        p = by_id.get(m['parent'])
        return f"Map {p['no']} side: {m['name']}" if p else m['name']
    return f"Map {m['no']}: {m['name']}"

def short_label(m):
    if m['sub']:
        p = by_id.get(m['parent'])
        return f"{p['no']}·{m['name']}" if p else m['name']
    return str(m['no'])

faced = defaultdict(list)     # card id -> [map labels] where an enemy fields it
facing = defaultdict(set)     # map id -> card ids its enemies field
ordered_maps = []
for m in mains:
    ordered_maps.append(m)
    ordered_maps.extend(subs_of[m['id']])
for m in ordered_maps:
    ids = set()
    for n in m['nodes']:
        for d in [n['deck']] + n['waves']:
            ids.update(d.keys())
        ids.update(n['leaders'])
    facing[m['id']] = ids
    for cid in ids:
        faced[cid].append(short_label(m))

cards = [c for c in data['cards'] if not c['test']]
cards_by_id = {c['id']: c for c in cards}

def source_label(c):
    s = c['src'] or {}
    k = s.get('kind')
    if c['token']: return 'Token (made by another card)'
    if k == 'map':
        m = by_id.get(s.get('id'))
        return f"Reward: {map_label(m)}" if m else 'Map reward'
    if k == 'pack':
        return PACK_POOL.get(int(s.get('tier') or 1), f"Pack pool {s.get('tier')}")
    if k == 'base': return 'Base (tutorial reward)' if s.get('tutorialReward') else 'Base'
    if k == 'smith': return 'Blacksmith (Forge)'
    if k == 'materia': return 'Materia (drop)'
    if k == 'design': return 'Design only (not in the game)'
    if k == 'secret': return 'Secret'
    if k: return k.capitalize()
    if c['rarity'] == 'starter' and not c['locked']: return 'Starter (everyone has it)'
    if c['id'] in faced: return 'Enemy only'
    return 'Unused'

def ability_lines(c):
    return [a for a in c['abilities'] if not a.startswith('🧬') and not a.startswith('⏳ Wait')]

card_rows = []
for c in sorted(cards, key=lambda c: (c['token'], c['name'].lower())):
    src = source_label(c)
    own = 'No' if (c['where'] == 'Not obtainable yet' or src.startswith(('Enemy only', 'Unused', 'Design', 'Token'))) else 'Yes'
    reward_map = ''
    s = c['src'] or {}
    if s.get('kind') == 'map' and s.get('id') in by_id:
        mm = by_id[s['id']]
        reward_map = by_id[mm['parent']]['no'] if mm['sub'] and mm['parent'] in by_id else mm['no']
    card_rows.append([
        c['name'], c['id'], 'Token' if c['token'] else KIND.get(c['kind'], c['kind'].capitalize()), RLABEL.get(c['rarity'], c['rarity']),
        c['cost'], c['wait'], c['attack'], c['health'], '\n'.join(ability_lines(c)), c['species'], c['types'],
        src, reward_map, ', '.join(dict.fromkeys(faced.get(c['id'], []))), own, c['where'], c['flavor']])
CARD_HEAD = ['Name', 'ID', 'Card kind', 'Rarity', 'Lumber cost', 'Wait', 'Attack', 'Health (defence)', 'Abilities',
             'Species', 'Types', 'Source', 'Reward map no.', 'Enemies field it on maps', 'Player can own', 'Where to get it (in-game text)', 'Flavour']

# ---------- map rows ----------
def node_code(m, key):
    for n in m['nodes']:
        if n['key'] == key: return n['code'] or key
    return key

def opens_text(m):
    if not m['sub']:
        prev = by_id.get(m['unlockAfter'])
        return f"After clearing Map {prev['no']} ({prev['name']})" if prev else 'From the start'
    e = m['entry'] or {}
    p = by_id.get(m['parent'])
    bits = []
    if e.get('after') and p: bits.append(f"after fight {node_code(p, e['after'])}")
    if e.get('afterMapCleared') and e['afterMapCleared'] in by_id: bits.append(f"after clearing {by_id[e['afterMapCleared']]['name']}")
    if e.get('afterDays'): bits.append(f"{e['afterDays']} days after")
    if e.get('quest'): bits.append(f"quest: {e['quest']}")
    return ('Opens ' + ', '.join(bits)) if bits else 'Opens with its map'

spots_at = defaultdict(list)
for sp in data['spots']:
    if sp['mapNumber'] is not None and sp['mapNumber'] <= 20:
        spots_at[sp['mapNumber']].append(f"{sp['icon']} {sp['name']}")
PACK_NAME = {p['id']: p['name'] for p in data['packs']}
free_pack = {int(k): PACK_NAME.get(v, v) for k, v in (data['freePacks'] or {}).items()}
merchant_at = {sp['mapNumber']: f"{sp['name']} sells the {PACK_NAME.get(sp['pack'], sp['pack'])}" for sp in data['spots'] if sp['merchant']}
wanderer = next((sp for sp in data['spots'] if sp['key'] == 'wanderer'), None)
if wanderer:
    merchant_at.setdefault(wanderer['mapNumber'], '')
    merchant_at[wanderer['mapNumber']] = '; '.join(x for x in [merchant_at[wanderer['mapNumber']], 'Wanderer’s Camp: Mystery Boosters (max 5 a visit, on days the trader appears)'] if x)

def level_text(levels):
    # roadEnemyLevel is 0 where enemy cards are not levelled up (before Map 5, and fights with their own setup)
    up = [l for l in levels if l]
    if not up: return 'Base level'
    lo, hi = min(up), max(up)
    rng = f"Lv {lo}" if lo == hi else f"Lv {lo}–{hi}"
    return rng + (' (some fights at base level)' if len(up) < len(levels) else '')

map_rows = []
for m in ordered_maps:
    fights = [n for n in m['nodes'] if n['kind'] != 'tutorial']
    bosses = [n['name'] for n in fights if 'boss' in n['kind']]
    ratings = [n['rating'] for n in fights if n['rating'] is not None]
    levels = [n['enemyLevel'] for n in fights if isinstance(n['enemyLevel'], (int, float))]
    castles = list(dict.fromkeys(n['castle'] for n in fights if n['castle']))
    rewards = [c['name'] for c in cards if (c['src'] or {}).get('kind') == 'map' and (c['src'] or {}).get('id') == m['id']]
    no = m['no'] if not m['sub'] else by_id[m['parent']]['no']
    map_rows.append([
        no, ('    ↳ ' if m['sub'] else '') + m['name'], 'Side map' if m['sub'] else 'Main map', m['id'],
        by_id[m['parent']]['name'] if m['sub'] else '', opens_text(m), len(fights), ', '.join(bosses),
        f"{min(ratings)}–{max(ratings)}" if ratings else '', level_text(levels),
        ', '.join(castles), len(rewards), ', '.join(rewards), len(facing[m['id']]),
        '' if m['sub'] else ', '.join(spots_at.get(m['no'], [])),
        '' if m['sub'] else (f"Boss first clear: 1 free {free_pack[m['no']]}" if m['no'] in free_pack else ''),
        '' if m['sub'] else merchant_at.get(m['no'], '')])
MAP_HEAD = ['Map no.', 'Map', 'Main / side', 'ID', 'Side map of', 'How it opens', 'Fights', 'Boss fights', 'Battle rating',
            'Enemy card level', 'Enemy castles', 'Reward cards (count)', 'Reward cards', 'Different enemy cards faced',
            'Features unlocked here', 'Free pack', 'Packs on sale from here']

# pack timeline
pool_sizes = defaultdict(int)
for c in cards:
    s = c['src'] or {}
    if s.get('kind') == 'pack' and not c['token']:
        pool_sizes[int(s.get('tier') or 1)] += 1
basic_pool = sum(pool_sizes[t] for t in (1, 11))
pack_rows = []
merch_by_pack = {sp['pack']: sp for sp in data['spots'] if sp['merchant']}
for p in data['packs']:
    pool = p['pool']
    content = f"{p['cards']} cards from the Mystery pool ({pool_sizes[99]} cards)" if pool == 99 else \
        f"{p['cards']} cards from the basic pools ({basic_pool} cards: pool 1 = {pool_sizes[1]}, Factions = {pool_sizes[11]})"
    sp = merch_by_pack.get(p['id'])
    if sp: sold = f"Map {sp['mapNumber']}: {sp['name']}"
    elif p['id'] == 'mystery' and wanderer: sold = f"Map {wanderer['mapNumber']}: Wanderer’s Camp (max 5 a visit, trader days only)"
    else: sold = ''
    free = ', '.join(f"Map {k}" for k, v in sorted(free_pack.items()) if v == p['name'])
    cost = ' + '.join(x for x in [f"{p['cost'].get('gold', 0)} Maple Leaves" if p['cost'].get('gold') else '', f"{p['cost'].get('gems', 0)} Gold Leaves" if p['cost'].get('gems') else ''] if x)
    pack_rows.append([p['name'], content, cost, sold, free])
PACK_HEAD = ['Pack', 'What is inside', 'Price', 'On sale from', 'One free from the boss of']

# ---------- effects ----------
def when_skill(desc, key):
    d = desc.lower()
    rules = [
        (('pitched',), 'When the card is pitched'),
        (('removal zone',), 'While in the Removal Zone (start of a round)'),
        (('on spawn', 'on play', 'on arrival', 'when it arrives', 'brings', 'arrives n turns ahead'), 'When played (on arrival)'),
        (('every round while ready, before melee',), 'Each round, before melee'),
        (('start of every round', 'start of each round', 'every round', 'each round'), 'Start of each round'),
        (('when this dies', 'perishes', 'any ally of yours dies'), 'When a unit dies'),
        (('kills', 'killing blow'), 'When it kills a unit'),
        (('when it becomes ready',), 'When it becomes Ready (Wait reaches 0)'),
        (('at night', 'by day', 'flow', 'ebb', 'rains', 'frost', 'moon'), 'Depends on day/night, weather or tide'),
        (('hits it', 'damages it', 'attackers take', 'hit taken', 'when first hit', 'incoming', 'takes n less', 'takes double', 'takes half', 'anything attacking', 'redirect', 'dodge'), 'When it is hit'),
        (('landed', 'on hit', 'every swing', 'when it attacks', 'as it attacks', 'after attacking', 'attacks twice', 'strikes', 'swings', 'lashes', 'throws on attack', 'its attacks', 'every attack', 'deals n', 'deals double', 'deals half', 'bonus damage', 'always attacks', 'hits for'), 'When it attacks or lands a hit'),
        (('enters play', 'another ally'), 'When another ally arrives'),
        (('whenever you gain',), 'When you gain a resource'),
        (('while', 'always', 'for every', 'can be your leader', 'counts for'), 'Always on while it is on the field'),
    ]
    for keys, label in rules:
        if any(k in d for k in keys): return label
    return 'Always on while it is on the field'

skill_users = defaultdict(list); trig_users = defaultdict(list); act_users = defaultdict(list); status_users = defaultdict(list)
for c in cards:
    if c['token']: continue
    for k in c['effectKeys']: skill_users[k].append(c['name'])
    for t in c['triggers']:
        on, do = t.split(':', 1)
        trig_users[on].append(c['name']); act_users[do].append(c['name'])
    text = ' '.join(ability_lines(c)).lower()
    for s in data['statuses']:
        if re.search(r'\b' + re.escape(s['name'].lower()[:5]), text): status_users[s['key']].append(c['name'])

def users(lst):
    lst = list(dict.fromkeys(lst))
    return len(lst), ', '.join(lst[:8]) + (f' … (+{len(lst) - 8})' if len(lst) > 8 else '')

effect_rows = []
for s in data['skills']:
    n, ex = users(skill_users[s['key']])
    desc = s.get('short') or (s['desc'] + (' (shown for N = 1)' if s.get('n1') else ''))
    effect_rows.append(['Skill', s['label'], s['key'], {'evergreen': 'Evergreen', 'passive': 'Passive'}.get(s['category'], s['category'].capitalize()), desc, when_skill(s['desc'], s['key']), n, ex])
for t in data['triggers']:
    n, ex = users(trig_users[t['key']])
    effect_rows.append(['Trigger (custom effects)', t['label'], t['key'], 'When a custom effect fires', t['desc'][0].upper() + t['desc'][1:], 'Fires ' + t['desc'], n, ex])
for a in data['actions']:
    n, ex = users(act_users[a['key']])
    effect_rows.append(['Action (custom effects)', a['label'], a['key'], 'What a custom effect does', a['desc'], 'When its trigger fires', n, ex])
for s in data['statuses']:
    n, ex = users(status_users[s['key']])
    effect_rows.append(['Status', s['name'], s['key'], 'Applied by skills and effects', f"A unit that is {s['name']} {s['text']}", 'While the status lasts', n, ex])
castle_users = defaultdict(list)
for m in ordered_maps:
    for nd in m['nodes']:
        if nd['castle']: castle_users[nd['castle']].append(short_label(m))
for c in data['castles']:
    where = list(dict.fromkeys(castle_users.get(c['name'], [])))
    effect_rows.append(['Castle passive', c['name'], c['id'], f"{c['hp']} HP castle", c['desc'], when_skill(c['desc'], c['id']) if c['desc'] and 'no passive' not in c['desc'].lower() else '—',
                        len(where), ('Enemy castle on maps: ' + ', '.join(where)) if where else ''])
EFFECT_HEAD = ['Effect type', 'Name', 'Key (in card data)', 'Group', 'What it does', 'When it happens', 'Cards using it', 'Example cards / where']

# ---------- workbook ----------
FONT = 'Arial'
head_fill = PatternFill('solid', start_color='2F4A2B')
band_fill = PatternFill('solid', start_color='EEF2E8')
sub_fill = PatternFill('solid', start_color='F7F9F3')
thin = Side(style='thin', color='C9D2C2')
wb = Workbook()

def write_table(ws, row0, head, rows, widths, wrap_cols=(), sub_col=None):
    for j, h in enumerate(head, 1):
        c = ws.cell(row=row0, column=j, value=h)
        c.font = Font(name=FONT, bold=True, color='FFFFFF', size=10); c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, r in enumerate(rows, row0 + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FONT, size=10, bold=(sub_col is not None and j == 2 and not str(r[1]).startswith('    ↳')))
            c.alignment = Alignment(wrap_text=(j in wrap_cols), vertical='top')
            c.border = Border(bottom=thin)
            if sub_col is not None and str(r[1]).startswith('    ↳'): c.fill = sub_fill
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    return row0 + len(rows)

def title(ws, text, note):
    ws['A1'] = text; ws['A1'].font = Font(name=FONT, bold=True, size=14)
    ws['A2'] = note; ws['A2'].font = Font(name=FONT, size=10, italic=True, color='555555')
    ws['A2'].alignment = Alignment(wrap_text=False)

ws = wb.active; ws.title = 'Cards'
own_n = sum(1 for r in card_rows if r[14] == 'Yes')
title(ws, 'Bramblewood Arena: cards', f"{len(card_rows)} cards ({own_n} a player can own). Snapshot of the game on {data.get('generated', '2026-10-12')}. "
      "Health is the card's defence. Abilities are the game's own card text. 'Enemies field it on maps' lists map numbers; 3·Name means a side map of Map 3.")
end = write_table(ws, 4, CARD_HEAD, card_rows, [24, 22, 11, 12, 8, 6, 7, 9, 70, 16, 24, 30, 9, 26, 8, 36, 50], wrap_cols=(9, 14, 16, 17))
ws.freeze_panes = 'C5'; ws.auto_filter.ref = f"A4:{get_column_letter(len(CARD_HEAD))}{end}"

ws = wb.create_sheet('Maps')
title(ws, 'Bramblewood Arena: maps', 'The road in play order (Map 0 to Map 20), each map followed by its side maps (↳). Battle ratings and enemy card levels are '
      'what the game computes for each fight. The card-pack timeline is under the table.')
end = write_table(ws, 4, MAP_HEAD, map_rows, [7, 30, 10, 6, 18, 34, 7, 30, 12, 10, 30, 9, 60, 11, 34, 26, 40], wrap_cols=(6, 8, 11, 13, 15, 16, 17), sub_col=2)
ws.freeze_panes = 'C5'
r = end + 3
ws.cell(row=r, column=1, value='When the card packs appear').font = Font(name=FONT, bold=True, size=13)
ws.cell(row=r + 1, column=1, value='Every basic pack (Sprout Pouch, Acorn Chest, Golden Bramble Case) draws from the same pools: pool 1 plus the Otter, Hummingbird '
        'and shared pools. Lumber-cost cards stay out of packs until pitching opens. See docs/pack-plan-2026-10-12.md for the proposed four-pack road structure.').font = Font(name=FONT, size=10, italic=True, color='555555')
end = write_table(ws, r + 3, PACK_HEAD, pack_rows, [], wrap_cols=(2, 4, 5))
for col, w in zip('ABCDE', [7, 30, 10, 6, 18]): pass  # widths stay those of the map table

ws = wb.create_sheet('Effects')
title(ws, 'Bramblewood Arena: effects', 'Every skill, custom-effect trigger and action, status and castle passive, with when it happens. '
      '"When it happens" for skills is read from the skill text; triggers say it themselves. N is the number set on each card.')
end = write_table(ws, 4, EFFECT_HEAD, effect_rows, [22, 22, 18, 22, 70, 34, 9, 60], wrap_cols=(5, 6, 8))
ws.freeze_panes = 'C5'; ws.auto_filter.ref = f"A4:H{end}"

wb.save(out)
for name, head, rows in (('cards', CARD_HEAD, card_rows), ('maps', MAP_HEAD, map_rows), ('effects', EFFECT_HEAD, effect_rows), ('packs', PACK_HEAD, pack_rows)):
    with open(f'{base}-{name}.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f); w.writerow(head); w.writerows(rows)
print(out, len(card_rows), 'cards', len(map_rows), 'maps', len(effect_rows), 'effects', len(pack_rows), 'packs')
