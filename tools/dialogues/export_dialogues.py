"""Export every NPC conversation in the game, with who says it and where it plays (2026-10-14, user: "Export out a list of
dialogues also - and where they come from").
Writes docs/dialogues.json (read by the hub) and docs/dialogues-<date>.xlsx (a sheet you can add your own rows to).
Needs a freshly assembled bramblewood-arena.html (python3 assemble_arena.py) or index.html."""
import json, os, re, sys, datetime, asyncio
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = open(os.path.join(ROOT, 'arena_app.js'), encoding='utf-8').read().split('\n')
PAGE = os.path.join(ROOT, 'bramblewood-arena.html' if os.path.exists(os.path.join(ROOT, 'bramblewood-arena.html')) else 'index.html')

async def grab():
    from playwright.async_api import async_playwright
    async with async_playwright() as p:
        b = await p.chromium.launch(); pg = await b.new_page()
        await pg.goto('file://' + PAGE); await pg.wait_for_timeout(2500)
        data = await pg.evaluate("""()=>{
          // the scripts built at run time: errands (one per template) and the leader lessons
          try{ (typeof ERRAND_TEMPLATES!=='undefined'?ERRAND_TEMPLATES:[]).forEach(t=>{ DIALOGUES['errand:'+t.id] = DIALOGUES['errand:'+t.id] || {lines:[{who:t.who, text:t.ask, options:[{label:'I’ll do it.'},{label:'Not now.'}]}, {who:t.who, text:'Thank you! I’ll pin it on the Notices so you don’t forget.'}, {who:t.who, text:'I’ll leave a note on the Quest Board, in case you change your mind.'}]}; }); }catch(e){}
          try{ DIALOGUES.lesson_leader_rule = leaderRuleScript(false); const hd = Object.values(getCardDefs()).find(x=> x.rarity === 'heroic' && x.art);
            DIALOGUES.lesson_leaders = leaderLessonScript(hd || {name:'Your new card', rarity:'heroic'}, false); }catch(e){}
          const mapName = id=>{ const m = (typeof CONQUEST_MAPS!=='undefined'?CONQUEST_MAPS:[]).find(x=> x.id === id); return m ? ('Map ' + mapNumberOf(id) + ' ' + (m.name||'')).trim() : id; };
          const defs = getCardDefs(), out = [];
          Object.keys(DIALOGUES).sort().forEach(id=>{ const sc = dialogueScript(id); (sc.lines||[]).forEach((l, i)=>{ const sp = speakerFor(l.who);
            const d = sp.card && defs[sp.card];
            out.push({id, n:i+1, who:l.who, name:sp.name, portrait: sp.art ? sp.art : (d && d.art ? 'card art: ' + d.name : 'emoji ' + (sp.icon||'💬')),
              text:l.text||'', replies:(l.options||[]).map(o=> o.label + (o.act ? ' ['+o.act+']' : '')).join(' | '), act:l.act||'', label:l.label||'', edited: !!(dialogueEdits[id] && dialogueEdits[id][i])}); }); });
          const spots = (typeof FEATURE_SPOTS!=='undefined'?FEATURE_SPOTS:[]).filter(s=> s.dialogue).map(s=> ({d:s.dialogue, where:'Map feature: ' + (s.label||s.name||s.key) + (s.map ? ' (' + mapName(s.map) + ')' : '')}));
          const curios = (typeof MAP_CURIOSITIES!=='undefined'?MAP_CURIOSITIES:[]).filter(s=> s.dialogue).map(s=> ({d:s.dialogue, where:'Map curiosity: ' + (s.name||s.label||s.key) + (s.map ? ' (' + mapName(s.map) + ')' : '')}));
          const nodes = []; (typeof CONQUEST_MAPS!=='undefined'?CONQUEST_MAPS:[]).forEach(m=> (m.nodes||[]).forEach(n=>{ if(n.dialogue) nodes.push({d:n.dialogue, where:'Before the fight ' + (n.code||n.key) + ' ' + (n.name||'') + ' (' + mapName(m.id) + ')'}); }));
          const speakers = Object.entries(DIALOGUE_SPEAKERS).map(([k, s])=> ({key:k, name:s.name, portrait: s.art || (s.card ? 'card art: ' + ((defs[s.card]||{}).name || s.card) : 'emoji ' + s.icon)}));
          return {lines:out, places: spots.concat(curios, nodes), speakers};
        }""")
        await b.close(); return data

def func_at(i):
    for j in range(i, -1, -1):
        m = re.match(r'\s*(?:async\s+)?function\s+(\w+)', SRC[j])
        if m: return m.group(1)
    return '?'

FRIENDLY = {'maybeLeaderRuleLesson':'First time you open the 👑 leader picker', 'maybeLeaderLesson':'After a fight, the first time you get a card that can lead (Heroic+); the example uses a Heroic card',
  'playFactionWarning':'A warning, in your faction’s voice', 'maybeOfferErrand':'After a Conquest win, sometimes (an errand offer)',
  'visitCuriosity':'Tapping a curiosity on the map', 'maybeShowCoachTip':'During a fight (tutorial and lessons)', 'renderMatchUI':'After finishing Tutorial Part 2', 'discardCardByUid':'In a fight, trying to send a 2nd card to the Graveyard in one turn',
  'pitchLessonTip':'In a fight, the first time a card needs Lumber', 'tutorialAllyRemarks':'Tutorial fights: remarks from your side', 'leavesRevealToMap':'The first time the Conquest map opens',
  'visitOldWell':'Tapping the old well on Map 0', 'renderCodex':'Opening the Codex once card families (phyla) unlock'}
SPECIFIC = {'both_peace':'After Tutorial Part 2, if you chose both tribes', 'training_wheels_off':'After Tutorial Part 2, the first time',
  'boss_banter':'The first time you start each boss fight (after the tutorial), in your side’s voices'}
FRIENDLY.update({'cobbleSay':'The Forge: Cobble the golem, when an upgrade is refused'})

# Dynamic ids (2026-10-18, audit P4: forge_cobble_* showed "(not found: maybe unused)" because cobbleSay builds the id as
# 'forge_cobble_' + key). Any playDialogue('prefix' + x) or playDialogue(`prefix${x}`) inside a function makes that function a
# wrapper for the prefix; a script whose id starts with the prefix is then found at each wrapper('suffix') call.
DYNAMIC = []   # (prefix, wrapper function name, line index)
for _i, _l in enumerate(SRC):
    for _m in re.finditer(r"""playDialogue\(\s*(?:(['"])([\w:@-]+)\1\s*\+|`([\w:@-]+)\$\{)""", _l):
        DYNAMIC.append((_m.group(2) or _m.group(3), func_at(_i), _i))

def dynamic_where(id, fallback=False):
    """Call sites of a wrapper that builds this id. With fallback (nothing else found), a wrapper with no literal call is named too."""
    out = []
    for prefix, fn, at in DYNAMIC:
        if not id.startswith(prefix) or fn == '?': continue
        suffix = id[len(prefix):]
        call = re.compile(r'\b' + re.escape(fn) + r"""\(\s*['"`]""" + re.escape(suffix) + r"""['"`]""")
        hits = [i for i, l in enumerate(SRC) if call.search(l)]
        for i in hits:
            f = func_at(i); out.append(FRIENDLY.get(fn, fn + '()') + ' from ' + f + f'()  · arena_app.js:{i+1}')
        if not hits and fallback: out.append(FRIENDLY.get(fn, fn + '()') + f' builds "{prefix}…" ids, but no {fn}(\'{suffix}\') call was found  · arena_app.js:{at+1}')
    return out

def where_for(id, places):
    if id in SPECIFIC: return [SPECIFIC[id]]
    out = [p['where'] for p in places if p['d'] == id]
    if id.startswith('errand:'): out.append('After a Conquest win, sometimes: an errand offer (' + id[7:] + ')')
    if id.startswith('warn:'): out.append('A warning, in your faction’s voice')
    out += dynamic_where(id)
    pat = re.compile(r"""['"`]""" + re.escape(id) + r"""['"`]""")
    start = next(i for i, l in enumerate(SRC) if l.startswith('const DIALOGUES = {'))
    end = next(i for i in range(start, len(SRC)) if SRC[i].startswith('};'))
    for i, l in enumerate(SRC):
        if start <= i <= end: continue
        if 'dialogue:' in l: continue  # a map spot's own definition: listed above by name
        if pat.search(l) and ('playDialogue' in l or 'DIALOGUES.' in l or 'dialogue' in l.lower()):
            f = func_at(i); out.append(FRIENDLY.get(f, f + '()') + f'  · arena_app.js:{i+1}')
    m = re.search(r'DIALOGUES\.' + re.escape(id) + r'\s*=', '\n'.join(SRC))
    if m:
        i = '\n'.join(SRC)[:m.start()].count('\n'); f = func_at(i); out.append(FRIENDLY.get(f, f + '()') + f'  · arena_app.js:{i+1}')
    if not out: out = dynamic_where(id, fallback=True)
    seen, res = set(), []
    for w in out:
        k = w.split('  · ')[0]
        if k not in seen: seen.add(k); res.append(w)
    return res or ['(not found: maybe unused)']

data = asyncio.run(grab())
ids = sorted(set(l['id'] for l in data['lines']))
where = {i: where_for(i, data['places']) for i in ids}
for l in data['lines']: l['where'] = ' ; '.join(where[l['id']]) if l['n'] == 1 else ''
os.makedirs(os.path.join(ROOT, 'docs'), exist_ok=True)
json.dump({'generated': datetime.date.today().isoformat(), 'speakers': data['speakers'], 'scripts': [{'id': i, 'where': where[i], 'lines': [l for l in data['lines'] if l['id'] == i]} for i in ids]},
          open(os.path.join(ROOT, 'docs', 'dialogues.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook(); ws = wb.active; ws.title = 'Dialogues'
F = 'Arial'; hdr = Font(name=F, bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='5B3A1E'); body = Font(name=F)
cols = [('Conversation', 26), ('Line', 6), ('Speaker key or card name', 22), ('Shown as', 28), ('Portrait', 26), ('Text', 70), ('Replies', 40), ('Where it plays', 60)]
ws.append([c for c, _ in cols])
for i, (_, w) in enumerate(cols): ws.column_dimensions[chr(65 + i)].width = w
for c in ws[1]: c.font = hdr; c.fill = fill
for l in data['lines']:
    ws.append([l['id'], l['n'], l['who'], l['name'], l['portrait'], l['text'] + (' ✎' if l['edited'] else ''), l['replies'], l['where']])
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = body; c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
sp = wb.create_sheet('Speakers'); sp.append(['Key', 'Name', 'Portrait'])
for c in sp[1]: c.font = hdr; c.fill = fill
for s in data['speakers']: sp.append([s['key'], s['name'], s['portrait']])
for col, w in zip('ABC', (16, 34, 34)): sp.column_dimensions[col].width = w
add = wb.create_sheet('Add your own')
rows = [['How to add your own lines'], ['Add rows on this sheet (one row per line) and send the file back; I put them in the game.'],
  ['Speaker: a key from the Speakers sheet, or any card’s name. A card speaks with its main card art.'],
  ['When: describe when it should play (a map spot, before a fight, after a win, the first time something happens...).'], [],
  ['Conversation', 'Line', 'Speaker key or card name', 'Text', 'Replies (separate with |)', 'When it should play'],
  ['example_bee_hello', 1, 'Bee Knight', 'Bzz! The hive sends its regards, champion.', 'Thanks! | Not now.', 'Example only: the first time you win a fight with Bee Knight in your deck']]
for r in rows: add.append(r)
add['A1'].font = Font(name=F, bold=True, size=13)
for c in add[6]: c.font = hdr; c.fill = fill
for c in add[7]: c.font = Font(name=F, italic=True, color='777777')
for col, w in zip('ABCDEF', (24, 6, 24, 60, 30, 50)): add.column_dimensions[col].width = w
out = os.path.join(ROOT, 'docs', f'dialogues-{datetime.date.today().isoformat()}.xlsx'); wb.save(out)
print(f'dialogues: {len(ids)} conversations, {len(data["lines"])} lines -> {out}')
