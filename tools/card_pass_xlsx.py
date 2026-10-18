#!/usr/bin/env python3
"""Card pass workbook (2026-10-19, package CB): docs/card-pass-2026-10-19.xlsx from docs/balance/card-pass-sheets.json
(written by tools/card_pass_docs_2026_10_19.js).

    node tools/card_pass_docs_2026_10_19.js && python3 tools/card_pass_xlsx.py

Sheets: Changed cards, Rarity, Sources, Sourceless, Odd sources kept, Left (owner calls), All cards, Curve, Decks.
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
data = json.load(open(os.path.join(ROOT, 'docs', 'balance', 'card-pass-sheets.json'), encoding='utf-8'))
HEAD = Font(bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='3A5A40')
FILL = {'IN': 'D1E7DD', 'green': 'D1E7DD', 'TOWARD': 'FFF3CD', 'EDGE': 'FFF3CD', 'FIT': 'E2E3E5', 'easy': 'CFE2FF', 'hard': 'F8D7DA'}
wb = Workbook(); wb.remove(wb.active)
for sh in data['sheets']:
    ws = wb.create_sheet(sh['name'][:31])
    ws.append(sh['header'])
    for c in ws[1]: c.font = HEAD; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='top')
    for r in sh['rows']:
        ws.append(['' if v is None else v for v in r])
        for v in r:
            if isinstance(v, str) and v in FILL:
                for c in ws[ws.max_row]: c.fill = PatternFill('solid', fgColor=FILL[v])
                break
    for i, h in enumerate(sh['header'], 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in sh['rows'][:400] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, width + 2))
    ws.freeze_panes = 'B2'
out = os.path.join(ROOT, 'docs', 'card-pass-2026-10-19.xlsx')
wb.save(out); print('wrote', out, len(data['sheets']), 'sheets')
