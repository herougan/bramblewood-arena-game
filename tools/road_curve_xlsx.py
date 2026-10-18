#!/usr/bin/env python3
"""Road curve workbook (2026-10-19, package R2): docs/road-curve-2026-10-19.xlsx from the pass's data files.

    python3 tools/road_curve_xlsx.py     reads docs/balance/skirmish-pass-r2-{before,after,fresh}.jsonl (tools/skirmish_pass.js measure
                                         --main), road-curve-nodes-{before,after}.json (every fight's CPU deck / Shield / Health before
                                         and after), road-curve-notes.json (why an outlier stays) and skirmish-pass-decks.json;
                                         writes docs/road-curve-2026-10-19.xlsx

Sheets: Fights (every main-map fight: map, position, kind, target and green band, before, after, what changed), Changes, Outliers
(after outside its green band, with the reason), Maps (green fights per map), Decks (the reference decks).
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
B = os.path.join(ROOT, 'docs', 'balance')
def jsonl(name):
    p = os.path.join(B, name)
    return [json.loads(l) for l in open(p, encoding='utf-8') if l.strip()] if os.path.exists(p) else []
def load(name, d=None):
    p = os.path.join(B, name)
    return json.load(open(p, encoding='utf-8')) if os.path.exists(p) else d
before = {(r['map'], r['key']): r for r in jsonl('skirmish-pass-r2-before.jsonl')}
after = {(r['map'], r['key']): r for r in jsonl('skirmish-pass-r2-after.jsonl')}
fresh = {(r['map'], r['key']): r for r in jsonl('skirmish-pass-r2-fresh.jsonl')}
nb, na = load('road-curve-nodes-before.json', {}), load('road-curve-nodes-after.json', {})
notes = load('road-curve-notes.json', {})
decks = load('skirmish-pass-decks.json', {'road': {}})
NAMES = {c['id']: c.get('name', c['id']) for c in json.load(open(os.path.join(ROOT, 'canonical', 'cards.json'), encoding='utf-8'))}
CPU = ['Basic', 'Smart', 'Smartest']
HEAD = Font(bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='3A5A40')
FILL = {'easy': PatternFill('solid', fgColor='CFE2FF'), 'hard': PatternFill('solid', fgColor='F8D7DA'), 'green': PatternFill('solid', fgColor='D1E7DD')}

def change_text(k):
    b, a = nb.get(k), na.get(k)
    if not b or not a: return ''
    out = []
    bd, ad = b.get('deck') or {}, a.get('deck') or {}
    for cid in sorted(set(bd) | set(ad), key=lambda c: NAMES.get(c, c)):
        d = ad.get(cid, 0) - bd.get(cid, 0)
        if d: out.append(f"{'+' if d > 0 else '−'}{abs(d)} {NAMES.get(cid, cid)}")
    if b.get('shield') != a.get('shield'): out.append(f"Shield {b.get('shield')} → {a.get('shield')}" + (' (override)' if a.get('shieldOverride') is not None else ' (rule)'))
    if b.get('hp') != a.get('hp'): out.append(f"Health {b.get('hp')} → {a.get('hp')}")
    return '; '.join(out)

def green(r, w):
    return 'easy' if w > r['hi'] + 1e-9 else 'hard' if w < r['lo'] - 1e-9 else 'green'

def sheet(wb, title, cols, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for c in ws[1]: c.font = HEAD; c.fill = HFILL; c.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    return ws

def order(k):
    r = before[k]; parts = r['code'].split('-')
    return (r['mapNo'], [int(x) if x.isdigit() else 0 for x in parts[1:]])

wb = Workbook(); wb.remove(wb.active)
cols = ['Code', 'Map', 'Map name', 'Fight', 'Kind', 'Role on the curve', 'Position', 'Target %', 'Green from', 'Green to', 'Your deck', 'Card Lv', 'CPU',
        'Win % before', 'Before', 'Win % after', 'After', 'Win % fresh seeds', 'Shield before', 'Shield after', 'Health before', 'Health after', 'What changed']
rows = []
for k in sorted(before, key=order):
    b = before[k]; a = after.get(k); f = fresh.get(k)
    pos = f"{b['pos']} of {b['of']}" if b.get('pos') else ('last (boss)' if b['role'] == 'boss' else '')
    aw = a['win'] if a else None
    rows.append([b['code'], b['mapNo'], b['mapName'], b['name'], b['kind'], b['role'], pos, b['target'], b['lo'], b['hi'], b['stage'], b['level'], CPU[b['cpu']],
                 b['win'], green(b, b['win']), aw, green(b, aw) if aw is not None else '', f['win'] if f else None,
                 b['shield'], a['shield'] if a else None, b['hp'], a['hp'] if a else None, change_text(b['map'] + ':' + b['key'])])
ws = sheet(wb, 'Fights', cols, rows, [8, 5, 22, 28, 10, 11, 10, 8, 8, 8, 9, 7, 9, 9, 8, 9, 8, 10, 8, 8, 8, 8, 80])
for r in ws.iter_rows(min_row=2):
    for idx in (14, 16):
        v = r[idx].value
        if v in FILL: r[idx].fill = FILL[v]
ch = [[r[0], r[2], r[3], r[5], r[7], r[13], r[15], r[17], r[22]] for r in rows if r[22]]
sheet(wb, 'Changes', ['Code', 'Map name', 'Fight', 'Role', 'Target %', 'Win % before', 'Win % after', 'Win % fresh seeds', 'What changed'], ch, [8, 22, 28, 10, 8, 10, 10, 10, 90])
outl = [[r[0], r[2], r[3], r[5], r[7], f"{r[8]:.0f}–{r[9]:.0f}", r[15], r[17], r[16], notes.get(r[0], '')] for r in rows if r[16] and r[16] != 'green']
sheet(wb, 'Outliers', ['Code', 'Map name', 'Fight', 'Role', 'Target %', 'Green', 'Win % after', 'Win % fresh seeds', 'After', 'Why it stays'], outl, [8, 22, 28, 10, 8, 8, 10, 10, 8, 90])
mrows = []
for no in sorted({r[1] for r in rows}):
    rs = [r for r in rows if r[1] == no]
    mrows.append([no, rs[0][2], len(rs), sum(1 for r in rs if r[14] == 'green'), sum(1 for r in rs if r[16] == 'green'),
                  sum(1 for r in rs if r[16] == 'easy'), sum(1 for r in rs if r[16] == 'hard'), sum(1 for r in rs if r[22])])
sheet(wb, 'Maps', ['Map', 'Name', 'Fights', 'Green before', 'Green after', 'Too easy after', 'Too hard after', 'Fights changed'], mrows, [5, 24, 8, 12, 12, 12, 12, 12])
drows = [['Starter', 'Map 0', 'the Otter starter (18 cards)']]
for no, d in sorted(decks.get('road', {}).items(), key=lambda x: int(x[0])):
    drows.append([f'Road {no}', f'Map {no}', ', '.join(f"{c['n']}x {c['name']}" for c in d['cards'])])
sheet(wb, 'Decks', ['Deck', 'Used on', 'Cards'], drows, [10, 10, 110])
dst = os.path.join(ROOT, 'docs', 'road-curve-2026-10-19.xlsx'); wb.save(dst)
print('wrote', dst, len(rows), 'fights,', sum(1 for r in rows if r[16] == 'green'), 'green after,', len(ch), 'changed,', len(outl), 'outliers')
