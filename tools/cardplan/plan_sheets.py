"""Card plan extra sheets (2026-10-14). Adds or rebuilds, in docs/bramblewood-card-plan-2026-10-14.xlsx:

* "Skill N values" (user: "In the sheet, let me choose the N values for EACH skill"): every skill key in the game's skill
  definitions (SKILL_DEFS in arena_app.js, read by tools/cardplan/extract_skills.js), its budget weight today
  (SB_DEFAULT_SKILLS, the Test Lab Stat budget) and a yellow "Your N" column. tools/cardplan/import_skill_n.py writes
  the filled values back into SB_DEFAULT_SKILLS. Rebuilding keeps any "Your N" and "Note" you already typed.
* "Factions & regions": the important skills of Base, the two factions and each region (counted from canonical/cards.json
  for the cards the Plan sheet puts there), each one's signature skills, and three proposed subfactions with one card each
  (design only, not in the game).
It also keeps the Plan and New cards sheets in step with the card renames and the four new design cards of 2026-10-14.
"""
import os, json, subprocess, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
BOOK = os.path.join(ROOT, 'docs', 'bramblewood-card-plan-2026-10-14.xlsx')
F = 'Arial'
HDR = Font(name=F, bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='3E5B3A')
TITLE = Font(name=F, bold=True, size=14); SUB = Font(name=F, italic=True, color='555555'); BODY = Font(name=F); BOLD = Font(name=F, bold=True)
YELLOW = PatternFill('solid', fgColor='FFF2CC'); GREY = PatternFill('solid', fgColor='EEEEEE'); BLUE_IN = Font(name=F, color='0000FF')
WRAP = Alignment(wrap_text=True, vertical='top')

data = json.loads(subprocess.check_output(['node', os.path.join(ROOT, 'tools', 'cardplan', 'extract_skills.js')]))
SKILLS, BUDGET, HALF_MISS, V3 = data['skills'], data['budget'], data.get('halfMiss', []), data.get('v3', {})
LABEL = {s['key']: s['label'] for s in SKILLS}; ICON = {s['key']: s['icon'] for s in SKILLS}
cards = json.load(open(os.path.join(ROOT, 'canonical', 'cards.json')))
CARD = {c['id']: c for c in cards}

wb = openpyxl.load_workbook(BOOK)

# ---- Plan / New cards: renames and the new design cards -------------------------------------------------------------
plan = wb['Plan']
for row in plan.iter_rows(min_row=2):
    cid = row[3].value
    if cid in CARD and cid in ('bee-tank', 'frog-imp', 'fly-imp', 'missingno'):
        row[2].value = CARD[cid]['name']
        if cid == 'missingno': row[15].value = 'Alien / irradiated / glitch. Map 21 (past the first world), not obtainable for now'
nc = wb['New cards']
have = {nc.cell(r, 1).value for r in range(5, nc.max_row + 1)}
STYLE = ('Pixel art, 32-bit era game sprite style, crisp hard edges, limited palette, dark 1-pixel outline, soft light from the upper left, '
         'warm muted woodland colours (moss green, bark brown, river teal, berry red) with one bright accent, cute expressive animal character '
         'with a slightly large head, clean readable silhouette, no text, no watermark, no signature. Card art: full-body character in a small '
         'scene, centred, slight low angle, simple painterly background. Subject: ')
NEW = [('rafflesia', 'Flower', 'a huge red rafflesia flower with white warts sprawled on the rainforest floor, a faint green stink haze rising from it, flies circling; background: dark rainforest undergrowth with hanging vines'),
       ('puffball-bombardier', 'Fungus', 'a round puffball mushroom with a stubby cap-helmet bursting a cloud of spores like a mortar shot; background: a glowing mushroom forest in a cave'),
       ('otter-angler', 'Otter', 'an otter sitting on a wooden weir with a willow fishing rod, a fish on the line, a little basket beside it; background: a calm river at dawn with reeds'),
       ('honeyjoe-brewer', 'Hummingbird', 'a hummingbird in a tiny apron hovering over a steaming acorn-cup of coffee, pouring honey with two bees helping; background: a flowering hive garden in the morning')]
for cid, animal, subject in NEW:
    if cid in have or cid not in CARD: continue
    c = CARD[cid]; e = c.get('effects') or {}
    sk = ', '.join(LABEL.get(k, 'Custom trigger' if k == 'triggers' else k) for k in e)
    nc.append([cid, c['name'], 'Design (not placed, 2026-10-14)', None, animal, c['rarity'].capitalize(), c['cost'], c['wait'], c['attack'], c['health'], sk, c['flavor'], STYLE + subject])
    for cell in nc[nc.max_row]: cell.font = BODY; cell.alignment = WRAP

# ---- Skill N values --------------------------------------------------------------------------------------------------
old = {}
if 'Skill N values' in wb.sheetnames:
    ws0 = wb['Skill N values']
    for r in ws0.iter_rows(min_row=1, values_only=True):
        if r and r[0] and isinstance(r[0], str) and r[0] in LABEL | {k: k for k in BUDGET}: old[r[0]] = (r[7], r[8])
    wb.remove(ws0)
ws = wb.create_sheet('Skill N values', index=wb.sheetnames.index('Skill curriculum') + 1)
use = collections.Counter(k for c in cards for k in (c.get('effects') or {}))
def counted(key, kind):
    w = BUDGET.get(key)
    if key in HALF_MISS: return f'Costs {w if w is not None else 0}: instead the card\'s whole score is ×{V3.get("missMult", 1.5)} (attacks miss it, or by it, about half the time)'
    if key == 'bounty': return 'Costs 0: Bounty pays both players (the one who plays it, and the other side when it dies)'
    if w == 'v': return f'By its number: {LABEL.get(key, key)} 2 costs 2 (at least 1; two-number values add up)'
    if isinstance(w, (int, float)): return f'Fixed: costs {w} whatever its number' if kind != 'flag' else f'Fixed: costs {w}'
    return 'Not in the table: costs 1' + ('' if kind == 'flag' else ', plus half its number (rounded down) when over 1')
ws.append(['Skill N values: what each skill costs in the stat budget'])
ws.append(['Rule (Test Lab ⚖️ Stat budget, v3, 2026-10-17): (2×Attack + Health + Skills) < N and (Attack × Health + Skills) < (N/3)², both must pass. Skills is the sum of the costs below. See the Budget v3 legend for N.'])
ws.append(['Fill the yellow "Your N" column: a number (the skill always costs that much) or v (it costs its own number, e.g. Poison 2 = 2). Leave it empty to keep the current weight. Then run: python3 tools/cardplan/import_skill_n.py'])
ws.append(['Kind: flag = on or off · number = it has a value (Poison 2) · trigger = a bundled When … do … skill. Blue text = values you type.'])
ws.append(['Example', '⚡', 'Quick', 'flag', 'boolean', 2, 'Fixed: costs 2', 3, 'Example only (this row is not read): Quick would cost 3 instead of 2.'])
ws.append([])
# Budget v3 legend (2026-10-17, user: "The N formula is changed. Now A*2 + B + Skills < N and A*B + Skills < (N/3)^2 And N is now
# 20 + L*2 + Rarity + Wait * 6 + 0.5 * Wait^2 + Lumber * 4 + 0.4 * Lumber^2. Base Rarity defined as 0, Common as 1, and onwards.")
# Replaces the v2 legend (N = 10 + 5×Wait + 4×Lumber + 0.2×map + level part + rarity part); the map term is gone.
v3g = lambda k, d: V3.get(k, d)
b = v3g('base', 20)
LEG = [('Budget v3', 'What N is made of (the same numbers as the game\'s Test Lab and card editor)'),
       ('Base', f'N starts at {b}'),
       ('Level', f'+{v3g("perLevel", 2)} per level (the Test Lab uses the level the card starts at: Base and quest cards 0, rewards 1; the card editor Level 1)'),
       ('Rarity', f'+{v3g("perRarity", 1)} per rarity step: Starter (Base) 0, Common 1, Uncommon 2, Rare 3, Very Rare 4, Super Rare 5, Epic 6, Heroic 7, Unique 8, Legendary 9, Mythic 10, Ancient 11, Divine 12, Void 13, Daemonic 14. Quest and special rarities count as the tier they stand for'),
       ('Wait', f'+{v3g("perWait", 6)}×Wait + {v3g("perWait2", 0.5)}×Wait²'),
       ('Lumber', f'+{v3g("perLumber", 4)}×Lumber + {v3g("perLumber2", 0.4)}×Lumber²'),
       ('Half-miss ×', f'Both scores are ×{v3g("missMult", 1.5)} for each of: {", ".join(HALF_MISS)} (they cost 0 themselves)'),
       ('Bounty', 'Costs 0: it pays the player who plays it and, when it dies, the other side'),
       ('Formula', f'N = {b} + {v3g("perLevel", 2)}×Level + Rarity + {v3g("perWait", 6)}×Wait + {v3g("perWait2", 0.5)}×Wait² + {v3g("perLumber", 4)}×Lumber + {v3g("perLumber2", 0.4)}×Lumber². Pass: 2A+H+S < N and A×H+S < (N/3)²')]
for k, v in LEG: ws.append([k, None, v])
LEG0 = ws.max_row - len(LEG) + 1
ws.append([])
H = ws.max_row + 1
ws.append(['Key', 'Icon', 'Skill', 'Kind', 'Value type', 'Current weight', 'How its value counts', 'Your N', 'Note', 'Cards using it'])
for c in ws[H]: c.font = HDR; c.fill = HFILL; c.alignment = WRAP
rows = [(s['key'], s['icon'], s['label'], s['kind'], s['valueKind']) for s in SKILLS]
rows += [(k, ICON.get(k, ''), k, 'number', 'number') for k in BUDGET if k not in LABEL]  # budget-table-only keys
for key, ic, lab, kind, vk in rows:
    w = BUDGET.get(key, '— (default rule)')
    note_default = '' if key in LABEL else 'Only in the budget table (no skill of this name in the editor list).'
    yn, nt = old.get(key, (None, None))
    ws.append([key, ic, lab, kind, vk, w, counted(key, kind), yn, nt if nt is not None else (note_default or None), use.get(key, 0)])
last = ws.max_row
for r in range(1, last + 1):
    for c in ws[r]:
        if r != H: c.font = BODY
ws['A1'].font = TITLE
for r in (2, 3, 4): ws.cell(r, 1).font = SUB
for r in range(LEG0, LEG0 + len(LEG)): ws.cell(r, 1).font = BOLD; ws.cell(r, 1).fill = GREY; ws.cell(r, 3).font = BODY
ws.cell(LEG0, 1).font = Font(name=F, bold=True, size=12); ws.cell(LEG0, 3).font = SUB
for c in ws[5]: c.font = Font(name=F, italic=True, color='777777'); c.fill = GREY
ws['H5'].font = Font(name=F, italic=True, color='0000FF'); ws['H5'].fill = YELLOW
for r in range(H + 1, last + 1):
    ws.cell(r, 8).fill = YELLOW; ws.cell(r, 8).font = BLUE_IN
    ws.cell(r, 7).alignment = WRAP; ws.cell(r, 9).alignment = WRAP
dv = DataValidation(type='custom', formula1=f'OR(H{H+1}="",H{H+1}="v",ISNUMBER(H{H+1}))', allow_blank=True,
                    errorTitle='Your N', error='Type a number, or v (the skill costs its own number).')
ws.add_data_validation(dv); dv.add(f'H{H+1}:H{last}')
for col, w in zip('ABCDEFGHIJ', (18, 7, 22, 9, 13, 13, 46, 10, 40, 12)): ws.column_dimensions[col].width = w
ws.freeze_panes = f'A{H+1}'
ws.auto_filter.ref = f'A{H}:J{last}'

# ---- Factions & regions ------------------------------------------------------------------------------------------------
if 'Factions & regions' in wb.sheetnames: wb.remove(wb['Factions & regions'])
fr = wb.create_sheet('Factions & regions', index=wb.sheetnames.index('Skill N values') + 1)
PH = [c.value for c in plan[1]]
prow = [dict(zip(PH, r)) for r in plan.iter_rows(min_row=2, values_only=True) if r[3]]
BIOMES = {'Swamp': ('Swamp',), 'Ice': ('Ice', 'Taiga'), 'Cave': ('Cave', 'Mushroom'), 'Ocean': ('Ocean', 'Coast', 'Mangrove', 'Sea cliffs', 'Beach'), 'Volcano': ('Volcano',),
          'River': ('River',), 'Desert': ('Desert',), 'Forest': ('Forest', 'Redwood', 'Briar', 'Hedgerow'), 'Mountain': ('Mountain', 'Painted cliffs'), 'Savanna': ('Savanna', 'Rainforest'),
          'Village and farm': ('Village', 'Plains', 'Castle')}  # the same regions as tools/cardplan/name_picker.py
def biome_of(b):
    b = str(b or '')
    for k, keys in BIOMES.items():
        if any(b.startswith(x) for x in keys): return k
groups = [('Base', lambda r: r['Group'] == 'Base'), ('Faction · Otters', lambda r: r['Group'] == 'Faction · Otters'),
          ('Faction · Hummingbirds', lambda r: r['Group'] == 'Faction · Hummingbirds')]
for b in BIOMES:
    groups.append((b, (lambda b: lambda r: biome_of(r['Biome']) == b and not str(r['Group']).startswith(('Retired', 'Faction', 'Base', 'Tutorial')))(b)))
def skill_keys(c):
    e = c.get('effects') or {}; out = []
    for k, v in e.items():
        if k == 'triggers':
            if v: out.append('triggers')
        elif k in LABEL: out.append(k)
    return out
name = lambda k: 'Custom trigger' if k == 'triggers' else LABEL.get(k, k)
icon = lambda k: '⚙️' if k == 'triggers' else ICON.get(k, '')
glob = collections.Counter(); total = 0
for c in cards:
    if c.get('token') or c.get('test'): continue
    total += 1; glob.update(set(skill_keys(c)))
summary, detail = [], []
for g, pred in groups:
    ids = [r['ID'] for r in prow if pred(r) and r['ID'] in CARD]
    wide = not ids and g in BIOMES  # e.g. River: every river card is a Base or faction card, so count those
    if wide: ids = [r['ID'] for r in prow if biome_of(r['Biome']) == g and not str(r['Group']).startswith('Retired') and r['ID'] in CARD]
    cnt = collections.Counter(k for i in ids for k in set(skill_keys(CARD[i])))
    n = len(ids)
    top = [k for k, _ in cnt.most_common() if k != 'triggers'][:6]
    # signature: used by 2+ cards here and far more common here than across all cards (lift), ties to the most used
    lift = {k: (v / n) / (glob[k] / total) for k, v in cnt.items() if v >= 2 and k != 'triggers'}
    sig = sorted(lift, key=lambda k: (-lift[k], -cnt[k]))[:2]
    summary.append((g + (' (its Base and faction cards)' if wide else ''), n, ', '.join(f'{icon(k)} {name(k)} ({cnt[k]})'.strip() for k in top) or '—', ', '.join(f'{icon(k)} {name(k)}'.strip() for k in sig) or '— (no skill stands out)'))
    for k, v in cnt.most_common(8): detail.append((g, icon(k), name(k), v, n, 'yes' if k in sig else ''))
fr.append(['Factions & regions: the skills each one leans on'])
fr.append(['Counted from canonical/cards.json for the cards the Plan sheet puts in each group (regions as in the name picker: biome fodder, map cards, later-world and spare cards of that biome). Signature = used by 2+ cards there and much more common there than across all cards.'])
fr.append([])
fr.append(['Group', 'Cards', 'Important skills (cards using it)', 'Signature skill(s)'])
H1 = fr.max_row
for c in fr[H1]: c.font = HDR; c.fill = HFILL
for s in summary:
    fr.append(list(s))
    for c in fr[fr.max_row]: c.font = BODY; c.alignment = WRAP
fr.append([])
fr.append(['Skill table', None, None, None, None, None, None])
fr.cell(fr.max_row, 1).font = Font(name=F, bold=True, size=12)
fr.append(['Group', 'Icon', 'Skill', 'Cards using it', 'Cards in group', 'Share', 'Signature?'])
H2 = fr.max_row
for c in fr[H2]: c.font = HDR; c.fill = HFILL
for g, ic, nm, v, n, sg in detail:
    fr.append([g, ic, nm, v, n, None, sg]); r = fr.max_row
    fr.cell(r, 6).value = f'=IF(E{r}=0,0,D{r}/E{r})'; fr.cell(r, 6).number_format = '0%'
    for c in fr[r]: c.font = BODY
fr.append([])
fr.append(['Proposed subfactions (design only, not in the game yet)'])
fr.cell(fr.max_row, 1).font = Font(name=F, bold=True, size=12)
fr.append(['Stats checked against the stat budget v3 (2026-10-17) at Level 1: N = 20 + 2×Level + Rarity + 6×Wait + 0.5×Wait² + 4×Lumber + 0.4×Lumber² (Common 1, Uncommon 2, Rare 3). Skills = the budget cost of the card\'s skills (Skill N values sheet).'])
fr.cell(fr.max_row, 1).font = SUB
fr.append(['Subfaction', 'Theme', 'Signature skills', 'Card', 'Animal', 'Rarity', 'Lumber', 'Wait', 'Attack', 'Health', 'Skills on the card', 'Skill cost', 'N', '2A+H+S < N', 'A×H+S < (N/3)²', 'Flavour'])
H3 = fr.max_row
for c in fr[H3]: c.font = HDR; c.fill = HFILL; c.alignment = WRAP
SUBF = [
    ('The Ashen Covenant (the Cultists)', 'A hooded cult that meets in burnt chapels and trades with the Devil: they give up their own units to call darker ones, and the ground they fight on goes bad.',
     '😈 Sacrifice, Inherent Darkness, Core Evil, Desecrate (Devilry archetype)', 'Ashen Acolyte', 'Rat (Rodent)', 'Uncommon', 0, 1, 2, 4, 'Sacrifice 1, Inherent Darkness 1', 2, 2,
     'Lights one black candle a night. Nobody in the burrow asks what the candles are for.'),
    ('The Great Herd (the Savannah cult)', 'The great beasts of the African savannah, lion, elephant, rhino, giraffe, buffalo and hippo, who worship the long grass and the rains. Slow, huge and loyal to each other.',
     '🐘 Overwhelm, 🫶 Kinship (Savanna), 😡 Rage, 🔨 Bludgeon', 'Old Tusker', 'Elephant (Megafauna)', 'Rare', 1, 2, 1, 10, 'Overwhelm, Kinship (Savanna +1/+1)', 2, 4,
     'The oldest elephant on the plain. Every herd that crosses it remembers the way to water.'),
    ('The Mycelium Choir', 'The mushroom folk under the forest, one root network that hums. They spread, they heal one another and they never really die: the spores always come back.',
     '🌱 Sprout, 🌿 Growth, Regeneration, 🦠 Festering', 'Lantern Morel', 'Mushroom (Elemental)', 'Uncommon', 0, 1, 1, 8, 'Lantern, Sprout 1', 2, 2,
     'Glows softly so the choir can find its way home. The little ones sprout wherever its light falls.'),
]
for (sf, theme, sig, cardn, animal, rar, lum, wait, a, h, sk, s, step, flav) in SUBF:
    fr.append([sf, theme, sig, cardn, animal, rar, lum, wait, a, h, sk, s, None, None, None, flav]); r = fr.max_row
    rank = {'Common': 1, 'Uncommon': 2, 'Rare': 3}.get(rar, 1)  # v3 (2026-10-17): the rarity rank, not v1's step
    fr.cell(r, 13).value = f'={v3g("base", 20)}+{v3g("perLevel", 2)}*1+{rank}+{v3g("perWait", 6)}*H{r}+{v3g("perWait2", 0.5)}*H{r}^2+{v3g("perLumber", 4)}*G{r}+{v3g("perLumber2", 0.4)}*G{r}^2'
    fr.cell(r, 14).value = f'=IF(2*I{r}+J{r}+L{r}<M{r},"✓ "&(2*I{r}+J{r}+L{r})&" < "&M{r},"✗ over")'
    fr.cell(r, 15).value = f'=IF(I{r}*J{r}+L{r}<(M{r}/3)^2,"✓ "&(I{r}*J{r}+L{r})&" < "&ROUND((M{r}/3)^2,1),"✗ over")'
    for c in fr[r]: c.font = BODY; c.alignment = WRAP
    fr.cell(r, 1).font = BOLD
for col, w in zip('ABCDEFGHIJKLMNOP', (30, 46, 40, 16, 18, 10, 8, 7, 8, 8, 30, 9, 7, 14, 12, 44)): fr.column_dimensions[col].width = w
fr.freeze_panes = f'A{H1+1}'

wb.save(BOOK)
print('Skill N values:', last - H, 'skills · Factions & regions:', len(summary), 'groups')
