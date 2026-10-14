"""Import the "Your N" values from the card plan's "Skill N values" sheet into the game's default stat-budget table
(SB_DEFAULT_SKILLS in arena_app.js). A number makes that skill always cost that much; v makes it cost its own number.
Empty cells are left alone, so an empty sheet changes nothing. Usage: python3 tools/cardplan/import_skill_n.py [--dry-run]

Note: a browser that saved its own costs in the Test Lab (⚖️ Stat budget > Skill costs) keeps those until it saves again."""
import os, re, sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
BOOK = os.path.join(ROOT, 'docs', 'bramblewood-card-plan-2026-10-14.xlsx')
APP = os.path.join(ROOT, 'arena_app.js')

ws = openpyxl.load_workbook(BOOK, data_only=True)['Skill N values']
head = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'Key')
cols = {ws.cell(head, c).value: c for c in range(1, ws.max_column + 1)}
want = {}
for r in range(head + 1, ws.max_row + 1):
    key, v = ws.cell(r, cols['Key']).value, ws.cell(r, cols['Your N']).value
    if not key or v is None or str(v).strip() == '': continue
    s = str(v).strip().lower()
    if s == 'v': want[key] = 'v'
    else:
        try: n = float(s)
        except ValueError: sys.exit(f'Row {r} ({key}): "Your N" must be a number or v, not {v!r}')
        want[key] = int(n) if n == int(n) else n

src = open(APP, encoding='utf-8').read()
m = re.search(r"^const SB_DEFAULT_SKILLS = \{(.*?)\};$", src, re.M)
if not m: sys.exit('SB_DEFAULT_SKILLS not found in arena_app.js')
cur = {}
for k, v in re.findall(r"(\w+)\s*:\s*('v'|-?\d+(?:\.\d+)?)", m.group(1)):
    cur[k] = 'v' if v == "'v'" else (float(v) if '.' in v else int(v))
changes = {k: (cur.get(k), v) for k, v in want.items() if cur.get(k) != v}
if not changes:
    print(f'No changes: {len(want)} value(s) filled in, all already match SB_DEFAULT_SKILLS ({len(cur)} entries).'); sys.exit(0)
for k, (a, b) in changes.items(): print(f'  {k}: {a if a is not None else "(default rule)"} -> {b}')
if '--dry-run' in sys.argv: print('Dry run: arena_app.js not written.'); sys.exit(0)
cur.update(want)
fmt = lambda v: "'v'" if v == 'v' else str(v)
line = 'const SB_DEFAULT_SKILLS = {' + ', '.join(f'{k}:{fmt(v)}' for k, v in cur.items()) + '};'
open(APP, 'w', encoding='utf-8').write(src[:m.start()] + line + src[m.end():])
print(f'Updated SB_DEFAULT_SKILLS in arena_app.js ({len(changes)} change(s)).')
