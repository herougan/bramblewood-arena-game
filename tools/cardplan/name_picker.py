"""Name picker (2026-10-14, user: "I'll just pick from your list of recommended NAMES - because some NAMES seem too advanced.
Then from NAMES we pick the effect etc. So I'll take 15, 15 + 15 Base, then 6 per biome").
Reads the card plan workbook and writes docs/name-picker-2026-10-14.xlsx: one sheet per group with every candidate name,
my ★ picks (simple, readable names first), and a ✓ column to fill in with a live count against the target."""
import os, re, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
plan = openpyxl.load_workbook(os.path.join(ROOT, 'docs', 'bramblewood-card-plan-2026-10-14.xlsx'))['Plan']
H = [c.value for c in plan[1]]; rows = [dict(zip(H, r)) for r in plan.iter_rows(min_row=2, values_only=True) if r[2]]
ADV = re.compile(r"\b(sovereign|warlord|archon|matriarch|patriarch|hierophant|exarch|paragon|harbinger|revenant|eldritch|leviathan|colossus|tyrant|regent|vanguard|sentinel|centurion|dominion|praetor|seraph|oracle|abyssal|primordial|unremembered|void|wraith|lich|ascendant|herald|scion|zealot|inquisitor|cardinal|vicar|templar|marshal|monarch|empress|emperor|baron|queen|king|lord|lady|duke)\b", re.I)
def level(n):
    w = len(n.split()); adv = bool(ADV.search(n)) or w >= 4 or len(n) > 22
    return 'Advanced' if adv else ('Simple' if w <= 2 and len(n) <= 16 and '-' not in n else 'Medium')
RANK = {'Simple': 0, 'Medium': 1, 'Advanced': 2}
RAR = ['starter', 'common', 'uncommon', 'rare', 'superrare', 'epic', 'heroic', 'unique', 'questunique', 'legendary', 'mythic', 'ancient']
def rar_i(r): r = str(r or '').lower().replace(' ', '').replace('-', ''); return RAR.index(r) if r in RAR else 5
BIOMES = {'Swamp': ('Swamp',), 'Ice': ('Ice', 'Taiga'), 'Cave': ('Cave', 'Mushroom'), 'Ocean': ('Ocean', 'Coast', 'Mangrove', 'Sea cliffs', 'Beach'), 'Volcano': ('Volcano',),
          'River': ('River',), 'Desert': ('Desert',), 'Forest': ('Forest', 'Redwood', 'Briar', 'Hedgerow'), 'Mountain': ('Mountain', 'Painted cliffs'), 'Savanna': ('Savanna', 'Rainforest'),
          'Village and farm': ('Village', 'Plains', 'Castle')}
def biome_of(b):
    b = str(b or '')
    for k, keys in BIOMES.items():
        if any(b.startswith(x) for x in keys): return k
    return None
NEW = {'Swamp': [('Vampire Bat', 'Proposed new (your idea): 3 Attack, drinks 3 Health back when it hits (Leech 3). Shared with Cave and Night maps')]}
SHARED = {'Swamp': 'Bog Leech already exists (2/5, Bleed 1) and can be the shared swamp leech.'}
groups = [('Base', lambda r: r['Group'] == 'Base', 15, 'Base cards everyone can get.'),
          ('Otters', lambda r: r['Group'] == 'Faction · Otters', 15, 'The otter faction.'),
          ('Hummingbirds', lambda r: r['Group'] == 'Faction · Hummingbirds', 15, 'The hummingbird faction.')]
for b in BIOMES:
    groups.append((b, (lambda b: lambda r: biome_of(r['Biome']) == b and not str(r['Group']).startswith(('Retired', 'Faction', 'Base', 'Tutorial')))(b), 6,
                   f'{b}: 6 unique cards; leftovers can be reused on nearby biomes.'))
wb = openpyxl.Workbook(); wb.remove(wb.active)
F = 'Arial'; hdr = Font(name=F, bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='4A5D23'); body = Font(name=F); star = Font(name=F, bold=True, color='8A5A00')
ov = wb.create_sheet('Overview')
ov.append(['Name picker']); ov['A1'].font = Font(name=F, bold=True, size=14)
ov.append(['Type ✓ in the Pick column of each sheet (or pick it from the drop-down). ★ marks my suggestion: simple, readable names first, then lower rarity. Name level flags names that read as advanced.'])
ov.append(['Once you have picked names, I set each card\'s skills and stats from the stat budget.']); ov.append([])
ov.append(['Group', 'Target', 'Picked', 'Left', 'Candidates'])
for c in ov[5]: c.font = hdr; c.fill = fill
for col, w in zip('ABCDE', (22, 9, 9, 9, 12)): ov.column_dimensions[col].width = w
ov.column_dimensions['A'].width = 22
for name, pred, target, note in groups:
    cand = [r for r in rows if pred(r)]
    for n, why in NEW.get(name, []): cand.append({'Name': n, 'ID': '(new)', 'Rarity': '', 'Attack': None, 'Health': None, 'Skills': '', 'Group': 'Proposed', 'Biome': name, 'Animal type': '', 'Notes': why})
    if not cand: print('empty', name); continue
    gp = lambda r: 0 if r['ID'] == '(new)' else 2 if str(r['Group']).startswith(('Later', 'Spare')) else 1  # your own ideas, then cards already placed
    cand.sort(key=lambda r: (RANK[level(r['Name'])], gp(r), rar_i(r['Rarity']), r['Name']))
    picks = set(id(r) for r in cand[:target])
    ws = wb.create_sheet(name[:31])
    ws.append([f'{name}: pick {target}', None, None, None, None, None, None, None, None]); ws['A1'].font = Font(name=F, bold=True, size=13)
    ws.append([note + (' ' + SHARED[name] if name in SHARED else '')]); ws.append([])
    ws.append(['Pick', '★', 'Name', 'Name level', 'Rarity', 'Stats', 'Skills now', 'Where it sits now', 'Animal type', 'ID'])
    for c in ws[4]: c.font = hdr; c.fill = fill
    for r in cand:
        st = f"{r['Attack']}/{r['Health']}" if r.get('Attack') is not None else ''
        ws.append(['', '★' if id(r) in picks else '', r['Name'], level(r['Name']), r.get('Rarity') or '', st, r.get('Skills') or '', r.get('Group') or '', r.get('Animal type') or '', r.get('ID') or ''])
    last = ws.max_row
    for row in ws.iter_rows(min_row=5, max_row=last):
        for c in row: c.font = body; c.alignment = Alignment(vertical='top', wrap_text=c.column in (7, 8))
        row[1].font = star
    dv = DataValidation(type='list', formula1='"✓"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'A5:A{last}')
    ws.conditional_formatting.add(f'A5:J{last}', FormulaRule(formula=['$A5="✓"'], fill=PatternFill('solid', fgColor='DCEFC8')))
    ws.conditional_formatting.add(f'D5:D{last}', FormulaRule(formula=['$D5="Advanced"'], font=Font(name=F, color='B03A2E')))
    for col, w in zip('ABCDEFGHIJ', (6, 4, 28, 11, 11, 7, 34, 30, 16, 26)): ws.column_dimensions[col].width = w
    ws.freeze_panes = 'C5'
    sh = name[:31].replace("'", "''")
    i = ov.max_row + 1
    ov.append([name, target, f"=COUNTIF('{sh}'!A5:A{last},\"✓\")", f'=B{i}-C{i}', len(cand)])
for row in ov.iter_rows(min_row=6):
    for c in row: c.font = body
out = os.path.join(ROOT, 'docs', 'name-picker-2026-10-14.xlsx'); wb.save(out); print(out)
