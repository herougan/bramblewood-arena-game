#!/usr/bin/env python3
"""Skill worth workbook (2026-10-19, package SW): docs/skill-worth-2026-10-19.xlsx.

    python3 tools/skill_worth_xlsx.py <sheet.json> <out.xlsx>

`node tools/skill_worth.js --report` writes the sheet rows ({sheets: [{name, head, rows}]}) from
docs/balance/skill-worth-2026-10-19.json and calls this; the rows are written as they are.
Sheets: Runs (one row per skill x N x field x baseline: win rates, worth by head-to-head and pool, S vs each opponent,
per-copy metrics and their change vs the vanilla unit, S today, suggested S), Entries (one row per skill-costs.json entry:
value today, measuredS, confidence, reason), Top gaps, Opponents (the calibrated reference decks).
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
data = json.load(open(src, encoding='utf-8'))
HEAD = Font(bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='3A5A40')
UNDER = PatternFill('solid', fgColor='D1E7DD'); OVER = PatternFill('solid', fgColor='F8D7DA'); LOW = Font(color='8A6D3B', italic=True)
wb = Workbook(); wb.remove(wb.active)
for sh in data['sheets']:
    ws = wb.create_sheet(sh['name'][:31])
    ws.append(sh['head'])
    for c in ws[1]:
        c.font = HEAD; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='top')
    gap_col = sh['head'].index('Gap') + 1 if 'Gap' in sh['head'] else None
    conf_col = sh['head'].index('Confidence') + 1 if 'Confidence' in sh['head'] else None
    for row in sh['rows']:
        ws.append(row)
        r = ws.max_row
        if gap_col:
            v = ws.cell(r, gap_col).value
            if isinstance(v, (int, float)) and abs(v) >= 1:
                ws.cell(r, gap_col).fill = UNDER if v > 0 else OVER
        if conf_col and ws.cell(r, conf_col).value == 'low':
            ws.cell(r, conf_col).font = LOW
    ws.freeze_panes = 'C2'
    for i, h in enumerate(sh['head'], 1):
        width = max([len(str(h))] + [len(str(x[i - 1])) for x in sh['rows'][:400] if i - 1 < len(x)])
        ws.column_dimensions[get_column_letter(i)].width = max(7, min(48, width + 2))
    ws.auto_filter.ref = ws.dimensions
wb.save(out)
print('wrote', out)
